"""Fabrique le fichier Excel qui devient le tableau Google Sheets des lieux."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LIGNES_VIDES = 500  # lignes prêtes à l'emploi (avec traduction automatique) sous les lieux existants

ENTETE = PatternFill("solid", fgColor="1D3244")
POLICE_ENTETE = Font(bold=True, color="FFFFFF")
AUTO = PatternFill("solid", fgColor="EEF4FA")  # colonnes remplies automatiquement
A_VERIFIER = PatternFill("solid", fgColor="FFF2CC")

LARGEURS_LIEUX = [30, 30, 24, 20, 22, 44, 70, 70, 60, 26, 26, 11, 11]
MODE_EMPLOI = [
    ("Mode d'emploi de la carte 3D Random Japan Place", True),
    ("", False),
    ("AJOUTER UN LIEU (à chaque nouveau TikTok)", True),
    ("1. Onglet « Lieux » : va sur la première ligne vide tout en bas.", False),
    ("2. Écris le nom en anglais (colonne A). Les noms français et japonais se remplissent tout seuls.", False),
    ("3. Choisis la catégorie dans le menu déroulant (colonne D).", False),
    ("4. Coordonnées GPS (colonne E) : dans Google Maps, fais un clic droit sur le lieu, clique sur les chiffres", False),
    ("   (ex. 34.4197, 131.0626) : ils sont copiés. Colle-les ici.", False),
    ("5. Colle le lien de ta vidéo TikTok (colonne F).", False),
    ("6. Écris la description en anglais (colonne G). Les traductions FR et JA se font toutes seules.", False),
    ("7. C'est tout ! Le lieu apparaît sur la carte en ~5 minutes.", False),
    ("", False),
    ("BON À SAVOIR", True),
    ("• Photo (colonne J) : laisse vide, la carte prend la miniature de la vidéo TikTok automatiquement.", False),
    ("  Tu peux aussi coller le lien d'une image (qui finit souvent par .jpg ou .png).", False),
    ("• Tu peux écrire par-dessus une traduction automatique si elle ne te plaît pas.", False),
    ("• Afficher ? = Non : le lieu est caché de la carte sans être effacé.", False),
    ("• À vérifier = Oui : lignes que Claude a remplies seul. Relis-les, corrige si besoin, puis efface le « Oui ».", False),
    ("", False),
    ("CATÉGORIES ET ICÔNES", True),
    ("• Onglet « Catégories » : ajoute une ligne pour créer une nouvelle catégorie.", False),
    ("• Colonne Icône : choisis un nom dans la liste (voir l'onglet « Icônes »), ou colle un emoji (ex. 🍜).", False),
    ("• Colonne Couleur : un code couleur, ex. #D63A2F (cherche « color picker » sur Google pour en trouver).", False),
    ("• Colonne Ordre : l'ordre d'affichage dans le menu de la carte (1 = en premier).", False),
    ("• Une catégorie sans aucun lieu n'apparaît pas sur la carte.", False),
    ("", False),
    ("⚠️ NE PAS CHANGER le nom des colonnes (ligne 1) ni le nom des onglets « Lieux » et « Catégories ».", True),
]


def _entete(ws, colonnes, largeurs):
    for i, (nom, large) in enumerate(zip(colonnes, largeurs), start=1):
        c = ws.cell(row=1, column=i, value=nom)
        c.fill, c.font = ENTETE, POLICE_ENTETE
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = large
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"


def fabriquer_excel(chemin, lieux, categories, icones, a_trier, colonnes_lieux, colonnes_categories):
    wb = Workbook()

    # --- Onglet Lieux -------------------------------------------------------
    ws = wb.active
    ws.title = "Lieux"
    _entete(ws, colonnes_lieux, LARGEURS_LIEUX)
    col = {nom: i + 1 for i, nom in enumerate(colonnes_lieux)}
    L = {nom: get_column_letter(i) for nom, i in col.items()}
    for r, lieu in enumerate(lieux, start=2):
        for nom, i in col.items():
            c = ws.cell(row=r, column=i, value=lieu.get(nom, "") or None)
            if nom.startswith("Description"):
                c.alignment = Alignment(wrap_text=False, vertical="top")
        if lieu.get("À vérifier") == "Oui":
            for i in col.values():
                ws.cell(row=r, column=i).fill = A_VERIFIER
    premiere_vide = len(lieux) + 2
    derniere = premiere_vide + LIGNES_VIDES - 1
    trad = [("Nom (FR)", "Nom (EN)", "fr"), ("Nom (日本語)", "Nom (EN)", "ja"),
            ("Description (FR)", "Description (EN)", "fr"), ("Description (日本語)", "Description (EN)", "ja")]
    for r in range(premiere_vide, derniere + 1):
        for cible, source, langue in trad:
            s = f"${L[source]}{r}"
            c = ws.cell(row=r, column=col[cible], value=f'=IF({s}="","",GOOGLETRANSLATE({s},"en","{langue}"))')
            c.fill = AUTO
        ws.cell(row=r, column=col["Afficher ?"], value="Oui")
    for r in range(2, derniere + 1):
        ws.cell(row=r, column=col["Coordonnées GPS"]).number_format = "@"

    dv_cat = DataValidation(type="list", formula1="=Catégories!$A$2:$A$200", allow_blank=True,
                            showErrorMessage=True, errorTitle="Catégorie inconnue",
                            error="Choisis une catégorie dans la liste (ou crée-la d'abord dans l'onglet Catégories).")
    dv_oui_non = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    dv_verif = DataValidation(type="list", formula1='"Oui"', allow_blank=True)
    for dv in (dv_cat, dv_oui_non, dv_verif):
        ws.add_data_validation(dv)
    dv_cat.add(f"{L['Catégorie']}2:{L['Catégorie']}{derniere}")
    dv_oui_non.add(f"{L['Afficher ?']}2:{L['Afficher ?']}{derniere}")
    dv_verif.add(f"{L['À vérifier']}2:{L['À vérifier']}{derniere}")

    # --- Onglet Catégories ----------------------------------------------------
    wc = wb.create_sheet("Catégories")
    _entete(wc, colonnes_categories, [24, 16, 12, 26, 26, 20, 8])
    for r, cat in enumerate(categories, start=2):
        for i, nom in enumerate(colonnes_categories, start=1):
            wc.cell(row=r, column=i, value=cat[nom])
        couleur = str(cat["Couleur"]).lstrip("#")
        wc.cell(row=r, column=3).fill = PatternFill("solid", fgColor=couleur)
        wc.cell(row=r, column=3).font = Font(color="FFFFFF", bold=True)
    dv_icone = DataValidation(type="list", formula1=f"=Icônes!$A$2:$A${len(icones) + 1}", allow_blank=True,
                              showErrorMessage=False)  # on peut aussi coller un emoji
    wc.add_data_validation(dv_icone)
    dv_icone.add("B2:B200")

    # --- Onglet Icônes --------------------------------------------------------
    wi = wb.create_sheet("Icônes")
    _entete(wi, ["Icône", "C'est quoi ?"], [16, 34])
    for r, (nom, desc) in enumerate(icones, start=2):
        wi.cell(row=r, column=1, value=nom)
        wi.cell(row=r, column=2, value=desc)
    wi.cell(row=len(icones) + 3, column=1, value="Pour voir les dessins : ouvre la page icones.html de ton site.")

    # --- Onglet À trier -------------------------------------------------------
    wt = wb.create_sheet("À trier")
    cols_t = ["Lien TikTok", "Légende TikTok", "Pourquoi", "Que faire ?"]
    _entete(wt, cols_t, [60, 90, 36, 40])
    for r, x in enumerate(a_trier, start=2):
        for i, nom in enumerate(cols_t, start=1):
            wt.cell(row=r, column=i, value=x.get(nom, ""))
    dv_faire = DataValidation(type="list", formula1='"Ignorer,Créer un lieu par endroit montré"', allow_blank=True)
    wt.add_data_validation(dv_faire)
    dv_faire.add(f"D2:D{len(a_trier) + 1}")

    # --- Onglet Mode d'emploi ---------------------------------------------------
    wm = wb.create_sheet("Mode d'emploi", 0)
    wm.column_dimensions["A"].width = 120
    for r, (texte, gras) in enumerate(MODE_EMPLOI, start=1):
        c = wm.cell(row=r, column=1, value=texte)
        c.font = Font(bold=gras, size=13 if r == 1 else 11)
    wb.active = 1  # s'ouvre sur l'onglet Lieux

    wb.save(chemin)
    return chemin
